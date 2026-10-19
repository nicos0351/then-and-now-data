import json
import sys
from collections import OrderedDict

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

rows = json.load(open(sys.argv[1]))
out_path = sys.argv[2]

FONT = 'Arial'
f = lambda **k: Font(name=FONT, **k)
HEAD_FILL = PatternFill('solid', fgColor='263238')
INPUT_FILL = PatternFill('solid', fgColor='FFF59D')
thin = Side(style='thin', color='D0D0D0')
BOX = Border(top=thin, bottom=thin, left=thin, right=thin)
STATUS_FILL = {
    'Keep': PatternFill('solid', fgColor='E8F5E9'),
    'Add': PatternFill('solid', fgColor='E3F2FD'),
    'Check': PatternFill('solid', fgColor='FFF3E0'),
    'Drop': PatternFill('solid', fgColor='F5F5F5'),
}

wb = Workbook()

# ---------------------------------------------------------------- Streets
ws = wb.active
ws.title = 'Streets'
cols = [
    ('Area', 26), ('Town / suburb', 18), ('Old name', 30), ('New name', 32), ('Suggested', 11),
    ('Founder decision', 16), ('Notes', 60), ('Source', 20), ('Latitude', 11), ('Longitude', 11),
    ('Map segments', 13), ('Road type', 13), ('See on map', 12),
]
for i, (name, width) in enumerate(cols, 1):
    c = ws.cell(row=1, column=i, value=name)
    c.font = f(bold=True, color='FFFFFF')
    c.fill = HEAD_FILL
    c.alignment = Alignment(vertical='center', wrap_text=True)
    ws.column_dimensions[get_column_letter(i)].width = width
ws.cell(row=1, column=6).fill = PatternFill('solid', fgColor='F9A825')
ws.cell(row=1, column=6).font = f(bold=True, color='000000')
ws.cell(row=1, column=6).comment = Comment('Fill in: Keep, Drop or Unsure. Blank means not reviewed yet.', 'Claude')
ws.cell(row=1, column=5).comment = Comment(
    "Keep = clear rename in the map data. Add = officially announced but missing from the map data. "
    "Check = plausible, confirm first. Drop = mapping detail or minor change.", 'Claude')
ws.cell(row=1, column=11).comment = Comment('How many OpenStreetMap pieces carry this rename; a rough guide to how long the road is.', 'Claude')
ws.row_dimensions[1].height = 30

for r_i, r in enumerate(rows, 2):
    values = [r['area'], r['town'], r['old'], r['new'], r['status'], None, r['note'], r['source'],
              r['lat'], r['lon'], r['segments'] or None, r['road_type'] or None, None]
    for c_i, v in enumerate(values, 1):
        c = ws.cell(row=r_i, column=c_i, value=v)
        c.font = f(size=10)
        c.border = BOX
        c.alignment = Alignment(vertical='top', wrap_text=c_i == 7)
    ws.cell(row=r_i, column=5).fill = STATUS_FILL[r['status']]
    ws.cell(row=r_i, column=6).fill = INPUT_FILL
    if r['lat'] is not None:
        link = ws.cell(row=r_i, column=13, value='Open map')
        link.hyperlink = f"https://www.openstreetmap.org/?mlat={r['lat']}&mlon={r['lon']}#map=17/{r['lat']}/{r['lon']}"
        link.font = f(size=10, color='1565C0', underline='single')
    for col in (9, 10):
        ws.cell(row=r_i, column=col).number_format = '0.00000'

last = len(rows) + 1
dv = DataValidation(type='list', formula1='"Keep,Drop,Unsure"', allow_blank=True)
dv.error = 'Choose Keep, Drop or Unsure'
dv.prompt = 'Keep, Drop or Unsure'
ws.add_data_validation(dv)
dv.add(f'F2:F{last}')
ws.conditional_formatting.add(f'F2:F{last}', CellIsRule(operator='equal', formula=['"Keep"'], fill=PatternFill('solid', fgColor='A5D6A7')))
ws.conditional_formatting.add(f'F2:F{last}', CellIsRule(operator='equal', formula=['"Drop"'], fill=PatternFill('solid', fgColor='E0E0E0')))
ws.freeze_panes = 'E2'
ws.auto_filter.ref = f'A1:M{last}'

# ---------------------------------------------------------------- Summary
s = wb.create_sheet('Summary', 0)
s.column_dimensions['A'].width = 32
for col in 'BCDEFG':
    s.column_dimensions[col].width = 14
s.sheet_view.showGridLines = False

s['A1'] = 'Street names add-on: renamed streets for review'
s['A1'].font = f(size=16, bold=True)
s['A2'] = 'Prepared 19 October 2026 from OpenStreetMap and official municipal announcements.'
s['A2'].font = f(size=10, italic=True, color='666666')

s['A4'] = 'How to review'
s['A4'].font = f(size=12, bold=True)
howto = [
    "On the Streets tab, fill in the yellow 'Founder decision' column with Keep, Drop or Unsure (pick from the list).",
    "'Suggested' is Claude's first sort: Keep, Add (announced but missing from the map data), Check (confirm first), Drop.",
    "'Open map' shows where the street is, so you can see the street's current name on the map.",
    'Example: Durban | West Street | Doctor Pixley Kaseme Street | Suggested Keep | Founder decision Keep.',
    'The table below updates as decisions are filled in.',
]
for i, t in enumerate(howto, 5):
    s.cell(row=i, column=1, value='•  ' + t).font = f(size=10)

top = 11
s.cell(row=top, column=1, value='By area').font = f(size=12, bold=True)
heads = ['Area', 'Keep', 'Add', 'Check', 'Drop', 'Founders: Keep', 'Not reviewed']
for i, h in enumerate(heads, 1):
    c = s.cell(row=top + 1, column=i, value=h)
    c.font = f(bold=True, color='FFFFFF')
    c.fill = HEAD_FILL
    c.alignment = Alignment(horizontal='center' if i > 1 else 'left', wrap_text=True)
s.row_dimensions[top + 1].height = 30

# Areas ordered by how much there is to show.
size = OrderedDict()
for r in rows:
    size.setdefault(r['area'], 0)
    if r['status'] in ('Keep', 'Add'):
        size[r['area']] += 1
areas = sorted(size, key=lambda a: (-size[a], a))
rng = lambda col: f"Streets!${col}$2:${col}${last}"
first = top + 2
for i, area in enumerate(areas):
    row = first + i
    s.cell(row=row, column=1, value=area).font = f(size=10)
    for j, status in enumerate(['Keep', 'Add', 'Check', 'Drop'], 2):
        s.cell(row=row, column=j, value=f'=COUNTIFS({rng("A")},$A{row},{rng("E")},"{status}")')
    s.cell(row=row, column=6, value=f'=COUNTIFS({rng("A")},$A{row},{rng("F")},"Keep")')
    s.cell(row=row, column=7, value=f'=COUNTIFS({rng("A")},$A{row},{rng("F")},"")')
    for j in range(1, 8):
        c = s.cell(row=row, column=j)
        c.border = BOX
        if j > 1:
            c.font = f(size=10)
            c.alignment = Alignment(horizontal='center')
end = first + len(areas) - 1
tot = end + 1
s.cell(row=tot, column=1, value='Total').font = f(bold=True)
for j in range(2, 8):
    col = get_column_letter(j)
    c = s.cell(row=tot, column=j, value=f'=SUM({col}{first}:{col}{end})')
    c.font = f(bold=True)
    c.alignment = Alignment(horizontal='center')
    c.border = Border(top=Side(style='medium'))

nt = tot + 2
s.cell(row=nt, column=1, value='What the research found').font = f(size=12, bold=True)
notes = [
    'OpenStreetMap has 5,087 South African street pieces tagged with an old name. Merged, they give 217 old/new pairs; 160 are clear renames.',
    'Durban (64), Pretoria (29), Pietermaritzburg (18), Cape Town (14) and Potchefstroom (11) have enough to make the add-on worthwhile.',
    "Pretoria matches the city's official 2012 list well: 26 of 28 lines are in the map data; the other 2 are added by hand.",
    "Johannesburg's 2014-15 inner-city renames (Bree, Jeppe, President, Noord, Sauer) are missing from the map data and are added by hand.",
    "Ekurhuleni's 44 renames (approved 2019) are almost all missing; 29 with a clear old street name are added by hand. Their positions must be set by hand.",
    'In 2011 a court reversed 9 of Durban\'s first-phase renames. The map data follows the current signs, but this is worth a check.',
    'Polokwane, Mbombela and most small towns: no street renames found (they were renamed as towns only).',
    'Map data © OpenStreetMap contributors (ODbL): the app must credit OpenStreetMap wherever these street names appear.',
]
for i, t in enumerate(notes, nt + 1):
    c = s.cell(row=i, column=1, value='•  ' + t)
    c.font = f(size=10)

# ---------------------------------------------------------------- Sources
src = wb.create_sheet('Sources')
src.column_dimensions['A'].width = 48
src.column_dimensions['B'].width = 100
src['A1'], src['B1'] = 'Source', 'Link'
for c in (src['A1'], src['B1']):
    c.font = f(bold=True, color='FFFFFF')
    c.fill = HEAD_FILL
sources = [
    ('OpenStreetMap (old_name tags), via Overpass API, October 2026', 'https://www.openstreetmap.org/copyright'),
    ('SAnews: Tshwane Metro to proceed with street name changes (2012)', 'https://www.sanews.gov.za/south-africa/tshwane-metro-proceed-street-name-changes'),
    ("The Heritage Portal: A Deeper Look at Johannesburg's New Street Names", 'https://www.theheritageportal.co.za/article/deeper-look-johannesburgs-new-street-names'),
    ('BusinessTech: These 44 roads in Ekurhuleni are getting a name change (Oct 2021)', 'https://businesstech.co.za/news/government/528874/these-44-roads-in-ekurhuleni-are-getting-a-name-change-what-you-need-to-know/'),
    ('IOL: Nine (Durban) street names to revert (Dec 2011)', 'https://iol.co.za/news/south-africa/kwazulu-natal/2011-12-01-nine-street-names-to-revert/'),
    ('Supreme Court of Appeal judgment on the Durban renames (2011)', 'https://lawlibrary.org.za/akn/za/judgment/zasca/2011/221'),
]
for i, (name, url) in enumerate(sources, 2):
    src.cell(row=i, column=1, value=name).font = f(size=10)
    c = src.cell(row=i, column=2, value=url)
    c.hyperlink = url
    c.font = f(size=10, color='1565C0', underline='single')

wb.save(out_path)
print('saved', out_path, 'rows', len(rows), 'areas', len(areas))
